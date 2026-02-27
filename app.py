"""
Sistema de Detecção de Impedimento Judicial
Interface web com Streamlit.

Uso:
    streamlit run app.py
"""

import io
import tempfile
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from detector import DetectorImpedimento, ResultadoProcesso


# ---------------------------------------------------------------------------
# Configuração da página
# ---------------------------------------------------------------------------

st.set_page_config(
    page_title="Detector de Impedimento",
    page_icon="⚖️",
    layout="wide",
)

st.title("⚖️ Sistema de Detecção de Impedimento Judicial")
st.markdown(
    """
    Faça o upload dos PDFs dos processos e informe os dados do juiz convocado.
    O sistema analisará cada processo e indicará se há **impedimento** de atuação
    no segundo grau por participação no primeiro grau.
    """
)


# ---------------------------------------------------------------------------
# Exportação para Excel
# ---------------------------------------------------------------------------

def _cor_nivel(nivel: str) -> str:
    return {"alto": "FF0000", "medio": "FFA500", "baixo": "FFD700"}.get(nivel, "FFFFFF")


def gerar_excel(resultados: list[ResultadoProcesso], nome_juiz: str) -> bytes:
    wb = Workbook()

    # --- Aba: Resumo ---
    ws_resumo = wb.active
    ws_resumo.title = "Resumo"

    cabecalho_resumo = [
        "Processo (arquivo)",
        "Resultado",
        "Nível de Confiança",
        "Vara Identificada",
        "Ocorrências",
        "Observações",
    ]
    ws_resumo.append(cabecalho_resumo)
    for cell in ws_resumo[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center")

    for r in resultados:
        if r.erro:
            linha = [r.arquivo, "ERRO", "", "", 0, r.erro]
            ws_resumo.append(linha)
            ws_resumo.cell(ws_resumo.max_row, 2).fill = PatternFill("solid", fgColor="AAAAAA")
        elif r.impedimento_detectado:
            linha = [
                r.arquivo,
                "IMPEDIMENTO DETECTADO",
                r.nivel_confianca.upper(),
                r.vara_identificada,
                len(r.ocorrencias),
                "",
            ]
            ws_resumo.append(linha)
            row = ws_resumo.max_row
            ws_resumo.cell(row, 2).fill = PatternFill("solid", fgColor=_cor_nivel(r.nivel_confianca))
            ws_resumo.cell(row, 2).font = Font(bold=True)
        else:
            linha = [r.arquivo, "Sem impedimento", "", r.vara_identificada, 0, ""]
            ws_resumo.append(linha)
            ws_resumo.cell(ws_resumo.max_row, 2).fill = PatternFill("solid", fgColor="C6EFCE")

    for col in ws_resumo.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws_resumo.column_dimensions[col[0].column_letter].width = min(max_len + 4, 60)

    # --- Aba: Ocorrências detalhadas ---
    ws_det = wb.create_sheet("Ocorrências Detalhadas")
    cabecalho_det = ["Processo", "Página", "Tipo de Ocorrência", "Contexto"]
    ws_det.append(cabecalho_det)
    for cell in ws_det[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")

    for r in resultados:
        for oc in r.ocorrencias:
            tipo_legivel = {
                "titulo": "Título de 1º grau",
                "assinatura": "Assinatura em decisão",
                "ato_jurisdicional": "Ato jurisdicional",
                "vara": "Vara do juiz",
            }.get(oc.tipo, oc.tipo)
            ws_det.append([r.arquivo, oc.pagina, tipo_legivel, oc.contexto])
            ws_det.cell(ws_det.max_row, 4).alignment = Alignment(wrap_text=True)

    ws_det.column_dimensions["A"].width = 35
    ws_det.column_dimensions["B"].width = 10
    ws_det.column_dimensions["C"].width = 25
    ws_det.column_dimensions["D"].width = 80

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Interface principal
# ---------------------------------------------------------------------------

with st.sidebar:
    st.header("⚙️ Configurações")

    nome_juiz = st.text_input(
        "Nome do juiz convocado",
        placeholder="Ex.: João da Silva Santos",
        help="Informe o nome completo ou as partes mais distintivas do nome.",
    )

    vara_juiz = st.text_input(
        "Vara do juiz (opcional)",
        placeholder="Ex.: 3ª Vara Cível de Campinas",
        help=(
            "Se informada, o sistema também pesquisará menções à vara "
            "como indicativo de impedimento."
        ),
    )

    st.markdown("---")
    st.markdown(
        "**Níveis de confiança:**\n"
        "- 🔴 **ALTO** – Nome + título de 1º grau e/ou assinatura\n"
        "- 🟠 **MÉDIO** – Nome + título ou vara confirmada\n"
        "- 🟡 **BAIXO** – Nome em ato jurisdicional isolado"
    )

st.header("📁 Upload dos Processos")

uploaded_files = st.file_uploader(
    "Selecione os PDFs dos processos",
    type=["pdf"],
    accept_multiple_files=True,
    help="Você pode selecionar múltiplos arquivos de uma só vez.",
)

if uploaded_files:
    st.info(f"{len(uploaded_files)} arquivo(s) selecionado(s).")

analisar = st.button("🔍 Analisar Processos", type="primary", disabled=not (uploaded_files and nome_juiz))

if not nome_juiz and uploaded_files:
    st.warning("Informe o nome do juiz convocado na barra lateral antes de analisar.")

# ---------------------------------------------------------------------------
# Processamento
# ---------------------------------------------------------------------------

if analisar and uploaded_files and nome_juiz:
    detector = DetectorImpedimento(nome_juiz=nome_juiz, vara_juiz=vara_juiz)
    resultados: list[ResultadoProcesso] = []

    progresso = st.progress(0, text="Iniciando análise…")
    status_texto = st.empty()

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)

        for i, uploaded in enumerate(uploaded_files):
            nome_arquivo = uploaded.name
            status_texto.text(f"Analisando: {nome_arquivo}")

            caminho = tmp / nome_arquivo
            caminho.write_bytes(uploaded.read())

            resultado = detector.analisar_pdf(caminho)
            resultados.append(resultado)

            progresso.progress((i + 1) / len(uploaded_files), text=f"{i+1}/{len(uploaded_files)} processos analisados")

    status_texto.empty()
    progresso.empty()

    # ------------------------------------------------------------------
    # Exibição dos resultados
    # ------------------------------------------------------------------

    st.header("📊 Resultados")

    col1, col2, col3, col4 = st.columns(4)
    total = len(resultados)
    com_impedimento = sum(1 for r in resultados if r.impedimento_detectado)
    erros = sum(1 for r in resultados if r.erro)
    sem_impedimento = total - com_impedimento - erros

    col1.metric("Total de processos", total)
    col2.metric("Com impedimento", com_impedimento, delta=None)
    col3.metric("Sem impedimento", sem_impedimento)
    col4.metric("Erros de leitura", erros)

    st.markdown("---")

    # Tabela resumo
    linhas = []
    for r in resultados:
        if r.erro:
            emoji = "❌"
            status = "Erro ao processar"
            confianca = ""
        elif r.impedimento_detectado:
            emoji = {"alto": "🔴", "medio": "🟠", "baixo": "🟡"}.get(r.nivel_confianca, "🔴")
            status = "IMPEDIMENTO DETECTADO"
            confianca = r.nivel_confianca.upper()
        else:
            emoji = "✅"
            status = "Sem impedimento"
            confianca = ""

        linhas.append({
            "": emoji,
            "Processo": r.arquivo,
            "Resultado": status,
            "Confiança": confianca,
            "Vara Identificada": r.vara_identificada,
            "Ocorrências": len(r.ocorrencias),
        })

    df = pd.DataFrame(linhas)
    st.dataframe(df, use_container_width=True, hide_index=True)

    # ------------------------------------------------------------------
    # Detalhamento por processo
    # ------------------------------------------------------------------

    processos_com_impedimento = [r for r in resultados if r.impedimento_detectado]

    if processos_com_impedimento:
        st.markdown("---")
        st.header("🔎 Detalhamento das ocorrências")

        for r in processos_com_impedimento:
            cor = {"alto": "red", "medio": "orange", "baixo": "goldenrod"}.get(r.nivel_confianca, "gray")
            with st.expander(f"📄 {r.arquivo} — :{cor}[{r.nivel_confianca.upper()}]"):
                if r.vara_identificada:
                    st.markdown(f"**Vara identificada no processo:** {r.vara_identificada}")

                for i, oc in enumerate(r.ocorrencias, 1):
                    tipo_legivel = {
                        "titulo": "Título de 1º grau",
                        "assinatura": "Assinatura em decisão/sentença",
                        "ato_jurisdicional": "Ato jurisdicional",
                        "vara": "Vara do juiz mencionada",
                    }.get(oc.tipo, oc.tipo)

                    st.markdown(f"**Ocorrência {i} — Página {oc.pagina} — {tipo_legivel}**")
                    st.code(oc.contexto, language=None)
                    st.divider()

    # ------------------------------------------------------------------
    # Download
    # ------------------------------------------------------------------

    st.markdown("---")
    st.header("⬇️ Exportar Relatório")

    excel_bytes = gerar_excel(resultados, nome_juiz)
    st.download_button(
        label="📥 Baixar relatório em Excel (.xlsx)",
        data=excel_bytes,
        file_name=f"impedimento_{nome_juiz.replace(' ', '_')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    # CSV simples
    csv_bytes = df.to_csv(index=False).encode("utf-8-sig")
    st.download_button(
        label="📥 Baixar resumo em CSV",
        data=csv_bytes,
        file_name=f"impedimento_{nome_juiz.replace(' ', '_')}.csv",
        mime="text/csv",
    )
